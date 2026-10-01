import openpyxl
import pytest

from generabase import salalm
from generabase.lectura import ArticuloSalalm
from generabase.salida import guardar_xlsx

CAD = {"N": (182, "ARTICULOS DE 6 MESES"), "K": (92, "ARTICULOS DE 3 MESES"), "J02": (85, "ARTICULOS DE 80 DIAS")}


def a(codigo, cajas=10.0, seccion="BEBIDAS", baja="", radubica="10012345", peso=8000, volumen=12648, unixcaja=6):
    return ArticuloSalalm(codigo, seccion, cajas, unixcaja, peso, volumen, seccion, baja, radubica)


def fila(articulo, maestro=None):
    return salalm.construir([articulo], maestro or {}, CAD)[0][0]


# 3.1
def test_peso_y_volumen_en_kg_y_litros():
    f = fila(a(273137, peso=8000, volumen=12648))
    assert (f.peso, f.volumen) == (8, 12.648)


def test_cruce_de_caducidad():
    f = fila(a(1), {1: ("S", "N")})
    assert (f.caduca, f.caducidad_codigo, f.caducidad, f.descripcion_caducidad, f.en_maestro) == ("S", "N", 182, "ARTICULOS DE 6 MESES", True)


def test_no_caduca_y_ausente_del_maestro():
    no_caduca, ausente = fila(a(1), {1: ("N", "")}), fila(a(2))
    assert (no_caduca.caduca, no_caduca.caducidad, no_caduca.en_maestro) == ("N", None, True)
    assert (ausente.caduca, ausente.caducidad, ausente.en_maestro) == ("", None, False)


def test_avisos_de_caducidad_incoherente():
    _, avisos = salalm.construir([a(1), a(2)], {1: ("S", "ZZ"), 2: ("S", "")}, CAD)
    assert len(avisos) == 2 and "'ZZ' no está" in avisos[0] and "MACADU = S" in avisos[1]


@pytest.mark.parametrize("radubica,ubic", [("24012404", 240), ("08004901", 80), ("   00000", None), ("00012345", 0), (6010221, 60), ("A1B00000", None)])
def test_ubicacion_tres_primeras_cifras(radubica, ubic):
    assert salalm.ubicacion(radubica) == ubic


@pytest.mark.parametrize("u,tipo", [(219, "ALIMENTACIÓN"), (220, "ALTA ROTACIÓN"), (290, "ALTA ROTACIÓN"), (291, "PASILLO ESTRECHO"),
                                    (399, "PASILLO ESTRECHO"), (400, "BAZAR ABAJO"), (599, "BAZAR ABAJO"), (600, "DROGUERÍA"), (1, "ALIMENTACIÓN")])
def test_tramos_de_ubicacion(u, tipo):
    assert salalm.tipo_ubicacion(u) == tipo


# 3.2
def test_exclusiones_y_recuento_por_regla():
    articulos = [
        a(1), a(2, seccion="suministros y servicios"), a(3, seccion=" CAMPAÑA FIDELIDAD "), a(4, baja="b"),
        a(5, radubica="   00000"), a(6, radubica="00099999"), a(7, radubica="35000000"), a(8, radubica="45000000"),
        a(9), a(10), a(11), a(12, seccion="SUMINISTROS Y SERVICIOS", baja="B"),
    ]
    maestro = {9: ("S", "J02"), 10: ("S", "K"), 11: ("N", "")}
    filas, _ = salalm.construir(articulos, maestro, CAD)
    quedan, excluidas = salalm.filtrar(filas)
    assert sorted(f.codigo for f in quedan) == [1, 10, 11]  # 92 días se queda; 85 fuera; sin caducidad se quedan
    assert excluidas == {"sección": 3, "baja": 1, "ubicación": 2, "tipo de ubicación": 2, "caducidad": 1}


def test_orden_por_cajas_desc_y_codigo():
    filas, _ = salalm.construir([a(5, cajas=3), a(2, cajas=7.5), a(9, cajas=7.5), a(1, cajas=100)], {}, CAD)
    quedan, _ = salalm.filtrar(filas)
    assert [(f.codigo, f.cajas) for f in quedan] == [(1, 100), (2, 7.5), (9, 7.5), (5, 3)]


# 3.3
def test_exportacion_salalm_filtrado(tmp_path):
    filas, _ = salalm.construir([a(273137, cajas=474.0), a(2, cajas=1.5, radubica="25000000")], {273137: ("S", "N")}, CAD)
    quedan, _ = salalm.filtrar(filas)
    destino = salalm.ruta_salalm_filtrado(tmp_path / "salalm.xlsx")
    guardar_xlsx(salalm.construir_libro(quedan), destino)
    assert destino.name == "SALALM FILTRADO.xlsx" and not list(tmp_path.glob(".~*"))
    ws = openpyxl.load_workbook(destino).active
    assert [c.value for c in ws[1]] == list(salalm.ENCABEZADOS)
    assert [c.value for c in ws[2]] == [273137, "BEBIDAS", 474, 6, 8, 12.648, "BEBIDAS", "S", "N", 182, "ARTICULOS DE 6 MESES", None, 100, "ALIMENTACIÓN"]
    assert [c.value for c in ws[3]][-2:] == [250, "ALTA ROTACIÓN"] and ws[3][7].value is None
    assert type(ws["C2"].value) is int and ws.freeze_panes == "A2" and ws["A1"].font.bold
