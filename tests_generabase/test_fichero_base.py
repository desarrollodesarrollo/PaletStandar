import openpyxl
import pytest

from generabase import fichero_base, salalm
from generabase.access import MovimientosUsados
from generabase.lectura import ArticuloSalalm
from generabase.modelo import ErrorValidacion
from generabase.salida import guardar_xlsx


def filas(*especificacion):
    """especificacion: (codigo, cajas[, radubica]) → filas de SALALM FILTRADO ya ordenadas."""
    articulos = [ArticuloSalalm(c, "SEC", cajas, 6, 6800, 8232, "SEC", "", r[0] if r else "10012345") for c, cajas, *r in especificacion]
    quedan, _ = salalm.filtrar(salalm.construir(articulos, {}, {})[0])
    return quedan


def usados(lineas):
    return MovimientosUsados(lineas=list(lineas), filas_leidas=len(lineas), filas_declaradas=len(lineas))


# 4.1
@pytest.mark.parametrize("valor", ["mil", "0", "-5", "2,5", "2.5", "", None])
def test_n_no_valido(valor):
    with pytest.raises(ErrorValidacion, match="nº de artículos debe ser un número entero mayor que 0"):
        fichero_base.parsear_n(valor)


def test_top_n_y_n_mayor_que_el_total():
    todas = filas((1, 5), (2, 50), (3, 20))
    elegidos, aviso = fichero_base.seleccionar(todas, 2)
    assert [f.codigo for f in elegidos] == [2, 3] and aviso is None
    elegidos, aviso = fichero_base.seleccionar(todas, 9000)
    assert len(elegidos) == 3 and "sólo tiene 3" in aviso


# 4.2
def test_bultos_y_lineas_por_articulo_y_tienda():
    elegidos = filas((2465665, 10))
    tiendas, agregado = fichero_base.agregar(elegidos, usados([(2465665, 389, 16 / 16), (2465665, 389, 8 / 16), (2465665, 81, 2.0)]))
    assert tiendas == [81, 389]
    assert agregado == {(2465665, 389): (1.5, 2), (2465665, 81): (2.0, 1)}


def test_tiendas_de_todos_los_movimientos_aunque_no_sean_de_los_n():
    elegidos = filas((1, 10))
    tiendas, agregado = fichero_base.agregar(elegidos, usados([(1, 82, 1.0), (999, 500, 3.0), (998, 81, 1.0)]))
    assert tiendas == [81, 82, 500] and agregado == {(1, 82): (1.0, 1)}


# 4.3
def test_estructura_hoja_todos(tmp_path):
    elegidos = filas((10, 80, "25000000"), (20, 30), (30, 5))
    mov = usados([(10, 81, 1.0), (10, 81, 0.5), (20, 82, 2.0), (99, 949, 1.0)])
    tiendas, agregado = fichero_base.agregar(elegidos, mov)
    libro = fichero_base.construir_libro(elegidos, tiendas, agregado)
    assert fichero_base.verificar_totales(libro, elegidos, mov) == (3.5, 3)
    destino = fichero_base.ruta_fichero_base(tmp_path / "salalm.xlsx")
    guardar_xlsx(libro, destino)
    assert destino.name == "FICHERO BASE.xlsx"
    ws = openpyxl.load_workbook(destino)["TODOS"]
    assert [c.value for c in ws[1]] == [None] * 5 + [81] * 3 + [82] * 3 + [949] * 3
    assert [c.value for c in ws[2]] == ["CODIGO", "TIPO", "PESO (Kg)", "Volumen (L)", "UNIxCAJA"] + ["BULTOS", "LINEAS", "ELECCIÓN"] * 3
    assert [c.value for c in ws[3]] == [10, "ALTA ROTACIÓN", 6.8, 8.232, 6, 1.5, 2, None, None, None, None, None, None, None]
    assert [c.value for c in ws[4]][5:11] == [None, None, None, 2, 1, None]
    assert type(ws["I4"].value) is int  # 2.0 bultos se escribe como entero
    assert [c.value for c in ws[5]][5:] == [None] * 9  # artículo sin movimientos
    assert ws.max_row == 5
    anchos = {l: ws.column_dimensions[l].width for l in "ABCDEFN"}
    assert anchos == {"A": 22.42578125, "B": 15.140625, "C": 13, "D": 13, "E": 15.140625, "F": 13, "N": 13}


def test_agregacion_manipulada_no_pasa_la_comprobacion():
    elegidos = filas((10, 80))
    mov = usados([(10, 81, 1.0), (10, 81, 0.5)])
    tiendas, agregado = fichero_base.agregar(elegidos, mov)
    agregado[(10, 81)] = (1.5, 1)
    libro = fichero_base.construir_libro(elegidos, tiendas, agregado)
    with pytest.raises(ErrorValidacion, match="LINEAS escritas 1 frente a 2"):
        fichero_base.verificar_totales(libro, elegidos, mov)
    agregado[(10, 81)] = (1.4, 2)
    with pytest.raises(ErrorValidacion, match="BULTOS escritos 1.400 frente a 1.500"):
        fichero_base.verificar_totales(fichero_base.construir_libro(elegidos, tiendas, agregado), elegidos, mov)
