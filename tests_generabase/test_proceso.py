import openpyxl
import pytest

from generabase import proceso
from generabase.modelo import ErrorEscritura, ErrorValidacion


def test_extremo_a_extremo(entradas):
    fases = []
    resumen = proceso.ejecutar(*entradas, "2", progreso=fases.append)
    carpeta = entradas[0].parent
    assert resumen.ruta_salalm == carpeta / "SALALM FILTRADO.xlsx" and resumen.ruta_base == carpeta / "FICHERO BASE.xlsx"
    assert fases[0] == "Leyendo SALALM…" and fases[-1] == "Guardando…" and any("Access" in f for f in fases)

    # SALALM FILTRADO: 104 baja, 105 sección, 106 sin ubicación, 107 pasillo estrecho, 108 caduca 85 días → fuera
    ws = openpyxl.load_workbook(resumen.ruta_salalm).active
    assert [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)] == [102, 101, 103]
    assert resumen.excluidas == {"sección": 1, "baja": 1, "ubicación": 1, "tipo de ubicación": 1, "caducidad": 1}
    assert (resumen.filas_salalm, resumen.filas_filtradas, resumen.sin_maestro) == (8, 3, 0)

    # FICHERO BASE con los 2 primeros (102 y 101); tiendas de todos los movimientos PI en rango
    wb = openpyxl.load_workbook(resumen.ruta_base)["TODOS"]
    assert [wb.cell(1, c).value for c in range(6, wb.max_column + 1, 3)] == [81, 349, 400, 500, 999]
    assert [wb.cell(r, 1).value for r in (3, 4)] == [102, 101] and wb.max_row == 4
    assert [wb.cell(3, c).value for c in (2, 3, 4, 5)] == ["ALTA ROTACIÓN", 8, 12.648, 6]
    assert [wb.cell(4, c).value for c in range(6, 12)] == [3, 2, None, None, None, None]  # 101 en 81: 2+1 cajas, 2 líneas
    assert [wb.cell(4, c).value for c in range(12, 15)] == [1.5, 1, None]  # 101 en 400
    assert [wb.cell(3, c).value for c in range(9, 12)] == [1, 1, None]  # 102 en 349
    assert (resumen.articulos_base, resumen.tiendas_base, resumen.bultos_base, resumen.lineas_base) == (2, 5, 5.5, 4)

    texto = resumen.texto()
    for esperado in ("SALALM FILTRADO: 3 artículos de 8", "FICHERO BASE: 2 artículos × 5 tiendas", "Filas leídas / declaradas: 8 / 8",
                     "Filas PI: 7", "PI en rangos de tienda: 6", "Suma de CAJAS: 7,50", "Fechas (MVFECI): 15/10/2025"):
        assert esperado in texto, esperado


def test_n_mayor_que_disponibles_avisa(entradas):
    resumen = proceso.ejecutar(*entradas, "50")
    assert resumen.articulos_base == 3 and "Se pidieron 50 artículos" in resumen.avisos[0]


def test_errores_no_generan_ficheros(entradas, tmp_path):
    with pytest.raises(ErrorValidacion, match="entero mayor que 0"):
        proceso.ejecutar(*entradas, "cero")
    salalm, maestro, caducidad, _ = entradas
    malo = tmp_path / "malo.accdb"
    malo.write_text("no es access")
    with pytest.raises(ErrorValidacion, match="no ser un Access"):
        proceso.ejecutar(salalm, maestro, caducidad, malo, "2")
    assert not (tmp_path / "SALALM FILTRADO.xlsx").exists() and not (tmp_path / "FICHERO BASE.xlsx").exists()


def test_reintento_sin_releer(entradas, monkeypatch):
    calculo = proceso.calcular(*entradas, "2")
    original = calculo.libro_base.save
    intentos = []

    def falla_una_vez(destino):
        intentos.append(destino)
        if len(intentos) == 1:
            raise PermissionError("abierto en Excel")
        original(destino)

    monkeypatch.setattr(calculo.libro_base, "save", falla_una_vez)
    monkeypatch.setattr(proceso.access, "leer_movimientos", lambda *a: pytest.fail("no debe releer el Access"))
    with pytest.raises(ErrorEscritura, match="FICHERO BASE.xlsx.*ciérralo y pulsa Reintentar"):
        calculo.guardar()
    assert not list(entradas[0].parent.glob(".~*"))
    resumen = calculo.guardar()
    assert resumen.ruta_base.exists() and resumen.ruta_salalm.exists() and len(intentos) == 2
