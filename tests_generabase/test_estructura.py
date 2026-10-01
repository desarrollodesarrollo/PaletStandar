import openpyxl

import generabase
from generabase import access, columnas, fichero_base, lectura, proceso, salalm

from .apoyo import FalsoAccessParser


def test_paquete_importable():
    assert generabase.__doc__
    for modulo in (access, columnas, fichero_base, lectura, proceso, salalm):
        assert modulo.__name__.startswith("generabase.")


def test_generadores_sinteticos(entradas):
    ruta_salalm, ruta_maestro, ruta_caducidad, ruta_access = entradas
    ws = openpyxl.load_workbook(ruta_salalm).active
    assert ws["C1"].value == "CODIGO" and ws["C2"].value == 101 and ws.max_row == 9
    assert openpyxl.load_workbook(ruta_maestro).active.title == "MSG001"
    assert openpyxl.load_workbook(ruta_caducidad).active["B2"].value == "K"
    db = FalsoAccessParser(ruta_access)
    assert "MOVIMIENTOS" in db.catalog
    assert len(db.parse_table("MOVIMIENTOS")["MVARTI"]) == db.get_table("MOVIMIENTOS").table_header["number_of_rows"] == 8
