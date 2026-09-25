from __future__ import annotations

import unicodedata
from fractions import Fraction
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .modelo import Articulo, DatosBase, DemandaTienda, ErrorValidacion, Medio

PRIMERA_COL_TIENDA = 5  # columna E
ENCABEZADOS_BLOQUE = ("BULTOS", "LINEAS", "ELECCION")
TIPOS_MEDIO = ("PALET", "CARRO")


def _texto_normalizado(valor) -> str:
    """Mayúsculas, sin espacios extremos y sin tildes."""
    texto = "" if valor is None else str(valor).strip().upper()
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def normalizar_codigo(valor):
    if valor is None:
        return None
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else valor
    texto = str(valor).strip()
    if not texto:
        return None
    if texto.isdigit():
        return int(texto)
    try:
        numero = float(texto.replace(",", "."))
    except ValueError:
        return texto
    return int(numero) if numero.is_integer() else texto


def a_fraccion(valor) -> Fraction | None:
    """Convierte un valor de celda a número exacto; None si no es numérico."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return Fraction(valor)
    if isinstance(valor, float):
        return Fraction(repr(valor))
    if isinstance(valor, str):
        texto = valor.strip().replace(",", ".")
        try:
            return Fraction(texto)
        except (ValueError, ZeroDivisionError):
            return None
    return None


def _vacio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _ultima_columna(filas_encabezado) -> int:
    ultima = 0
    for fila in filas_encabezado:
        for i, valor in enumerate(fila, start=1):
            if not _vacio(valor):
                ultima = max(ultima, i)
    return ultima


def extraer_datos_base(ws: Worksheet) -> DatosBase:
    filas = list(ws.iter_rows(values_only=True))
    if len(filas) < 2:
        raise ErrorValidacion("El fichero base no tiene las dos filas de encabezado.")
    fila1, fila2 = filas[0], filas[1]
    ultima = _ultima_columna((fila1, fila2))
    n_cols_tienda = ultima - (PRIMERA_COL_TIENDA - 1)
    if n_cols_tienda <= 0:
        raise ErrorValidacion("El fichero base no tiene columnas de tiendas a partir de la columna E.")
    if n_cols_tienda % 3:
        raise ErrorValidacion(
            f"El nº de columnas de tiendas desde la columna E hasta la {get_column_letter(ultima)} "
            f"({n_cols_tienda}) no es múltiplo de 3 (BULTOS, LINEAS, ELECCIÓN)."
        )

    def celda(fila, col):
        return fila[col - 1] if col - 1 < len(fila) else None

    if _texto_normalizado(celda(fila2, 4)) != "UNIXCAJA":
        raise ErrorValidacion(
            f"Celda D2: la columna D del fichero base debe tener el encabezado 'UNIxCAJA' (unidades por caja) "
            f"y tiene '{'' if celda(fila2, 4) is None else celda(fila2, 4)}'. ¿Es un fichero base antiguo?"
        )

    bloques = []
    vistas = {}
    for col in range(PRIMERA_COL_TIENDA, ultima + 1, 3):
        for k, esperado in enumerate(ENCABEZADOS_BLOQUE):
            valor = celda(fila2, col + k)
            if _texto_normalizado(valor) != esperado:
                nombre = "ELECCIÓN" if esperado == "ELECCION" else esperado
                raise ErrorValidacion(
                    f"Columna {get_column_letter(col + k)}, fila 2: se esperaba el encabezado "
                    f"'{nombre}' y hay '{'' if valor is None else valor}'."
                )
        codigos = [normalizar_codigo(celda(fila1, col + k)) for k in range(3)]
        if codigos[0] is None or len(set(codigos)) != 1:
            raise ErrorValidacion(
                f"Columnas {get_column_letter(col)}-{get_column_letter(col + 2)}, fila 1: las tres "
                f"columnas del bloque deben tener el mismo código de tienda (hay {codigos})."
            )
        tienda = codigos[0]
        if tienda in vistas:
            raise ErrorValidacion(
                f"La tienda {tienda} está repetida (columnas {get_column_letter(vistas[tienda])} "
                f"y {get_column_letter(col)})."
            )
        vistas[tienda] = col
        bloques.append((tienda, col))

    articulos: list[Articulo] = []
    avisos: list[str] = []
    filas_datos = []
    for n_fila, fila in enumerate(filas[2:], start=3):
        codigo = normalizar_codigo(celda(fila, 1))
        if codigo is None:
            continue
        peso, volumen = a_fraccion(celda(fila, 2)), a_fraccion(celda(fila, 3))
        valido = peso is not None and volumen is not None and peso > 0 and volumen > 0
        if not valido:
            avisos.append(
                f"Artículo {codigo} (fila {n_fila}): peso o volumen vacío o no válido; no se seleccionará."
            )
        valor_unidades = celda(fila, 4)
        unidades = a_fraccion(valor_unidades)
        if unidades is None or unidades.denominator != 1 or unidades < 1:
            raise ErrorValidacion(
                f"Fila {n_fila}, columna D (artículo {codigo}): UNIxCAJA debe ser un número entero mayor o igual "
                f"que 1 y es '{'' if valor_unidades is None else valor_unidades}'."
            )
        articulos.append(
            Articulo(
                fila=n_fila,
                codigo=codigo,
                peso=peso or Fraction(0),
                volumen=volumen or Fraction(0),
                valido=valido,
                unidades_caja=int(unidades),
            )
        )
        filas_datos.append((n_fila, fila))

    tiendas = []
    for tienda, col in bloques:
        bultos, lineas = [], []
        for n_fila, fila in filas_datos:
            for destino, c, nombre in ((bultos, col, "BULTOS"), (lineas, col + 1, "LINEAS")):
                valor = celda(fila, c)
                if _vacio(valor):
                    destino.append(Fraction(0))
                    continue
                numero = a_fraccion(valor)
                if numero is None or numero < 0:
                    raise ErrorValidacion(
                        f"Fila {n_fila}, columna {get_column_letter(c)} (tienda {tienda}): el valor de "
                        f"{nombre} debe ser un número mayor o igual que 0 y es '{valor}'."
                    )
                destino.append(numero)
        tiendas.append(DemandaTienda(tienda=tienda, col_eleccion=col + 2, bultos=bultos, lineas=lineas))

    return DatosBase(articulos=articulos, tiendas=tiendas, avisos=avisos)


def leer_fichero_base(ruta: Path) -> tuple[Workbook, DatosBase]:
    wb = _abrir(ruta, "fichero base")
    return wb, extraer_datos_base(wb.worksheets[0])


def extraer_medios(ws: Worksheet, tiendas_necesarias) -> dict:
    necesarias = set(tiendas_necesarias)
    filas: dict = {}
    for n_fila, fila in enumerate(ws.iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        fila = tuple(fila) + (None,) * (5 - len(fila))
        tienda = normalizar_codigo(fila[0])
        if tienda not in necesarias:
            continue
        if tienda in filas and filas[tienda][1] != fila[1:]:
            raise ErrorValidacion(
                f"La tienda {tienda} aparece con datos distintos en las filas {filas[tienda][0]} y {n_fila} "
                "de la tabla de tiendas y medios."
            )
        filas.setdefault(tienda, (n_fila, fila[1:]))

    faltan = [t for t in tiendas_necesarias if t not in filas]
    if faltan:
        raise ErrorValidacion(
            f"Faltan {len(faltan)} tiendas del fichero base en la tabla de tiendas y medios: "
            + ", ".join(str(t) for t in faltan)
        )

    medios = {}
    for tienda in tiendas_necesarias:
        n_fila, (tipo, peso, volumen, densidad) = filas[tienda]
        tipo_norm = _texto_normalizado(tipo)
        if tipo_norm not in TIPOS_MEDIO:
            raise ErrorValidacion(
                f"Tienda {tienda} (fila {n_fila} de la tabla): tipo de medio '{tipo}' no válido; "
                "debe ser PALET o CARRO."
            )
        numeros = []
        for nombre, valor in (("PESO MAX", peso), ("VOLUMEN MÁX", volumen), ("DENSIDAD", densidad)):
            numero = a_fraccion(valor)
            if numero is None or numero <= 0:
                raise ErrorValidacion(
                    f"Tienda {tienda} (fila {n_fila} de la tabla): {nombre} debe ser un número mayor que 0 "
                    f"y es '{'' if valor is None else valor}'."
                )
            numeros.append(numero)
        medios[tienda] = Medio(tipo_norm, *numeros)
    return medios


def leer_tabla_medios(ruta: Path, tiendas_necesarias) -> dict:
    wb = _abrir(ruta, "tabla de tiendas y medios", solo_lectura=True)
    try:
        return extraer_medios(wb.worksheets[0], tiendas_necesarias)
    finally:
        wb.close()


def _abrir(ruta: Path, descripcion: str, solo_lectura: bool = False) -> Workbook:
    try:
        return openpyxl.load_workbook(ruta, read_only=solo_lectura)
    except FileNotFoundError:
        raise ErrorValidacion(f"No se encuentra el {descripcion}: {ruta}") from None
    except PermissionError:
        raise ErrorValidacion(f"No se puede abrir el {descripcion} (¿está abierto en otro programa?): {ruta}") from None
    except Exception as exc:
        raise ErrorValidacion(f"El {descripcion} no es un Excel .xlsx válido: {ruta} ({exc})") from None
