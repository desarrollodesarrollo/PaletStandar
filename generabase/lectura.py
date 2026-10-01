"""Lectura y validación de SALALM, MAESTRO y MAESTRO DE CADUCIDAD."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .columnas import codigo_texto, entero, indice, normalizar, numero, texto
from .modelo import ErrorValidacion

COLUMNAS_SALALM = ("CODIGO", "CAJAS", "DESCRIPC", "DESSEC", "BAJA", "UNIXCAJA", "RADUBICA", "PESO", "VOLUMEN")
COLUMNAS_MAESTRO = ("MAARTI", "MACADU", "MACLCA")
COLUMNAS_CADUCIDAD = ("C3CADU", "C3UTIL", "C3COME")


@dataclass(frozen=True)
class ArticuloSalalm:
    codigo: int
    descripcion: str
    cajas: float
    unixcaja: float
    peso: float
    volumen: float
    seccion: str
    baja: str
    radubica: object


def _filas(ruta: Path, nombre: str):
    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except FileNotFoundError:
        raise ErrorValidacion(f"No se encuentra el fichero {nombre}: {ruta}") from None
    except PermissionError:
        raise ErrorValidacion(f"No se puede abrir el fichero {nombre} (¿está abierto en otro programa?): {ruta}") from None
    except Exception as exc:
        raise ErrorValidacion(f"El fichero {nombre} no es un Excel .xlsx válido: {ruta} ({exc})") from None
    try:
        filas = list(libro.worksheets[0].iter_rows(values_only=True))
    finally:
        libro.close()
    if not filas:
        raise ErrorValidacion(f"El fichero {nombre} está vacío.")
    return filas


def _vacia(fila) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in fila)


def _celda(fila, pos):
    return fila[pos] if pos < len(fila) else None


def leer_salalm(ruta: Path) -> tuple[list[ArticuloSalalm], list[str]]:
    nombre = "SALALM"
    filas = _filas(ruta, nombre)
    col = indice(filas[0], COLUMNAS_SALALM, nombre)
    articulos, avisos, vistos = [], [], {}
    for n_fila, fila in enumerate(filas[1:], start=2):
        if _vacia(fila):
            continue
        valor_codigo = _celda(fila, col["CODIGO"])
        codigo = entero(valor_codigo)
        if codigo is None:
            avisos.append(f"SALALM fila {n_fila}: CODIGO '{texto(valor_codigo)}' no es un número; se ignora la fila.")
            continue
        if codigo in vistos:
            raise ErrorValidacion(f"SALALM: el CODIGO {codigo} está repetido en las filas {vistos[codigo]} y {n_fila}.")
        vistos[codigo] = n_fila
        numeros = {}
        for c in ("CAJAS", "UNIXCAJA", "PESO", "VOLUMEN"):
            valor = _celda(fila, col[c])
            numeros[c] = numero(valor)
            if numeros[c] is None:
                raise ErrorValidacion(
                    f"SALALM fila {n_fila}, columna {get_column_letter(col[c] + 1)} ({c}): "
                    f"el valor '{texto(valor)}' no es un número."
                )
        articulos.append(
            ArticuloSalalm(
                codigo=codigo,
                descripcion=texto(_celda(fila, col["DESCRIPC"])),
                cajas=numeros["CAJAS"],
                unixcaja=numeros["UNIXCAJA"],
                peso=numeros["PESO"],
                volumen=numeros["VOLUMEN"],
                seccion=texto(_celda(fila, col["DESSEC"])),
                baja=texto(_celda(fila, col["BAJA"])),
                radubica=_celda(fila, col["RADUBICA"]),
            )
        )
    return articulos, avisos


def leer_maestro(ruta: Path) -> tuple[dict[int, tuple[str, str]], list[str]]:
    """{MAARTI: (MACADU, MACLCA)}; un MAARTI repetido conserva la primera aparición."""
    nombre = "MAESTRO"
    filas = _filas(ruta, nombre)
    col = indice(filas[0], COLUMNAS_MAESTRO, nombre)
    maestro, avisos, primera = {}, [], {}
    for n_fila, fila in enumerate(filas[1:], start=2):
        if _vacia(fila):
            continue
        codigo = entero(_celda(fila, col["MAARTI"]))
        if codigo is None:
            avisos.append(f"MAESTRO fila {n_fila}: MAARTI '{texto(_celda(fila, col['MAARTI']))}' no es un número; se ignora.")
            continue
        if codigo in maestro:
            avisos.append(f"MAESTRO: MAARTI {codigo} repetido (filas {primera[codigo]} y {n_fila}); se usa la fila {primera[codigo]}.")
            continue
        primera[codigo] = n_fila
        maestro[codigo] = (normalizar(_celda(fila, col["MACADU"])), codigo_texto(_celda(fila, col["MACLCA"])))
    return maestro, avisos


def leer_caducidades(ruta: Path) -> dict[str, tuple[float, str]]:
    """{C3CADU: (C3UTIL, C3COME)}."""
    nombre = "MAESTRO DE CADUCIDAD"
    filas = _filas(ruta, nombre)
    col = indice(filas[0], COLUMNAS_CADUCIDAD, nombre)
    caducidades, primera = {}, {}
    for n_fila, fila in enumerate(filas[1:], start=2):
        if _vacia(fila):
            continue
        clave = codigo_texto(_celda(fila, col["C3CADU"]))
        if not clave:
            continue
        if clave in caducidades:
            raise ErrorValidacion(
                f"MAESTRO DE CADUCIDAD: el código C3CADU '{clave}' está repetido en las filas {primera[clave]} y {n_fila}."
            )
        valor = _celda(fila, col["C3UTIL"])
        dias = numero(valor)
        if dias is None:
            raise ErrorValidacion(
                f"MAESTRO DE CADUCIDAD fila {n_fila}, columna {get_column_letter(col['C3UTIL'] + 1)} (C3UTIL): "
                f"el valor '{texto(valor)}' no es un número."
            )
        primera[clave] = n_fila
        caducidades[clave] = (int(dias) if float(dias).is_integer() else dias, texto(_celda(fila, col["C3COME"])))
    return caducidades
