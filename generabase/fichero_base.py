"""FICHERO BASE: los N artículos con más cajas y sus bultos/líneas por tienda, con la estructura de la plantilla."""

from __future__ import annotations

import math
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .access import MovimientosUsados
from .modelo import NOMBRE_FICHERO_BASE, ErrorValidacion
from .salalm import FilaSalalm

HOJA = "TODOS"
ENCABEZADOS_ARTICULO = ("CODIGO", "TIPO", "PESO (Kg)", "Volumen (L)", "UNIxCAJA")
ENCABEZADOS_TIENDA = ("BULTOS", "LINEAS", "ELECCIÓN")
PRIMERA_COL_TIENDA = len(ENCABEZADOS_ARTICULO) + 1  # F
ANCHOS = {"A": 22.42578125, "B": 15.140625, "C": 13.0, "D": 13.0, "E": 15.140625}
ANCHO_TIENDA = 13.0
_ENTERO = re.compile(r"^\+?\d+$")


def parsear_n(valor) -> int:
    texto = "" if valor is None else str(valor).strip()
    if not _ENTERO.match(texto) or int(texto) < 1:
        raise ErrorValidacion("El nº de artículos debe ser un número entero mayor que 0.")
    return int(texto)


def seleccionar(filas: list[FilaSalalm], n: int) -> tuple[list[FilaSalalm], str | None]:
    if n > len(filas):
        return list(filas), f"Se pidieron {n} artículos pero SALALM FILTRADO sólo tiene {len(filas)}; se usan todos."
    return list(filas[:n]), None


def agregar(articulos: list[FilaSalalm], usados: MovimientosUsados) -> tuple[list[int], dict[tuple[int, int], tuple[float, int]]]:
    """Tiendas (todas las de los movimientos usados, ascendentes) y {(artículo, tienda): (bultos, líneas)}."""
    codigos = {a.codigo for a in articulos}
    cajas = defaultdict(list)
    for articulo, tienda, c in usados.lineas:
        if articulo in codigos:
            cajas[(articulo, tienda)].append(c)
    return usados.tiendas, {k: (math.fsum(v), len(v)) for k, v in cajas.items()}


def _numero(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


def construir_libro(articulos: list[FilaSalalm], tiendas: list[int], agregado: dict) -> Workbook:
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA
    fila1 = [None] * len(ENCABEZADOS_ARTICULO)
    fila2 = list(ENCABEZADOS_ARTICULO)
    for t in tiendas:
        fila1 += [t] * 3
        fila2 += ENCABEZADOS_TIENDA
    hoja.append(fila1)
    hoja.append(fila2)
    for a in articulos:
        fila = [a.codigo, a.tipo_ubicacion, _numero(a.peso), _numero(a.volumen), _numero(a.unixcaja)]
        for t in tiendas:
            bultos, lineas = agregado.get((a.codigo, t), (None, None))
            fila += [None if bultos is None else _numero(bultos), lineas, None]
        hoja.append(fila)
    for letra, ancho in ANCHOS.items():
        hoja.column_dimensions[letra].width = ancho
    for c in range(PRIMERA_COL_TIENDA, PRIMERA_COL_TIENDA + 3 * len(tiendas)):
        hoja.column_dimensions[get_column_letter(c)].width = ANCHO_TIENDA
    return libro


def verificar_totales(libro: Workbook, articulos: list[FilaSalalm], usados: MovimientosUsados) -> tuple[float, int]:
    """Recalcula desde los movimientos los totales esperados y los compara con lo escrito en la hoja."""
    codigos = {a.codigo for a in articulos}
    esperado_bultos = math.fsum(c for a, _, c in usados.lineas if a in codigos)
    esperado_lineas = sum(1 for a, _, _ in usados.lineas if a in codigos)
    hoja = libro[HOJA]
    bultos, lineas = [], 0
    for fila in hoja.iter_rows(min_row=3, min_col=PRIMERA_COL_TIENDA, values_only=True):
        for k in range(0, len(fila), 3):
            if fila[k] is not None:
                bultos.append(fila[k])
            if fila[k + 1] is not None:
                lineas += fila[k + 1]
    escrito_bultos = math.fsum(bultos)
    if lineas != esperado_lineas or not math.isclose(escrito_bultos, esperado_bultos, rel_tol=1e-9, abs_tol=1e-9):
        raise ErrorValidacion(
            "La comprobación final del FICHERO BASE ha fallado; no se ha guardado ningún fichero. "
            f"BULTOS escritos {escrito_bultos:.3f} frente a {esperado_bultos:.3f} esperados; "
            f"LINEAS escritas {lineas} frente a {esperado_lineas} esperadas."
        )
    return esperado_bultos, esperado_lineas


def ruta_fichero_base(ruta_salalm: Path) -> Path:
    return Path(ruta_salalm).parent / NOMBRE_FICHERO_BASE
