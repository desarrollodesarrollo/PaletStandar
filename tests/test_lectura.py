from fractions import Fraction

import openpyxl
import pytest

from pruebaestandar.lectura import extraer_datos_base, extraer_medios, leer_fichero_base, leer_tabla_medios
from pruebaestandar.modelo import ErrorValidacion, Medio

from .conftest import CARRO, PALET, crear_base, crear_tabla, datos_aleatorios


def hoja(ruta):
    return openpyxl.load_workbook(ruta).active


# 2.1
def test_medio_carro_duplica_capacidad():
    carro = Medio("CARRO", Fraction(475), Fraction(600), Fraction(475, 600))
    palet = Medio("PALET", Fraction(875), Fraction(1300), Fraction(875, 1300))
    assert (carro.n_medios, carro.peso_max, carro.volumen_max) == (2, 950, 1200)
    assert (palet.n_medios, palet.peso_max, palet.volumen_max) == (1, 875, 1300)
    assert carro.densidad == Fraction(475, 600)


# 2.2
@pytest.mark.parametrize("n_art,n_tiendas", [(130, 180), (977, 3)])
def test_lectura_dimensiones(tmp_path, n_art, n_tiendas):
    articulos, tiendas, _ = datos_aleatorios(n_art, n_tiendas, semilla=1)
    _, datos = leer_fichero_base(crear_base(tmp_path / "b.xlsx", articulos, tiendas))
    assert len(datos.articulos) == n_art
    assert [d.tienda for d in datos.tiendas] == [t[0] for t in tiendas]
    assert datos.tiendas[-1].col_eleccion == 4 + 3 * n_tiendas
    i, (bultos, lineas) = next(iter(tiendas[0][1].items()))
    assert datos.tiendas[0].bultos[i] == bultos and datos.tiendas[0].lineas[i] == lineas


def test_celdas_vacias_son_cero(tmp_path):
    ruta = crear_base(tmp_path / "b.xlsx", [(1, 2, 4), (2, 3, 5)], [(10081, {0: (7, 2)})])
    _, datos = leer_fichero_base(ruta)
    assert datos.tiendas[0].bultos == [7, 0]
    assert datos.tiendas[0].lineas == [2, 0]


def test_codigos_normalizados_y_filas_sin_codigo(tmp_path):
    ruta = crear_base(tmp_path / "b.xlsx", [(1, 2, 4), (None, 3, 5), ("0003", 1, 2)], [("10081", {0: (1, 1)})])
    ws = hoja(ruta)
    ws.cell(1, 6).value = 10081.0
    datos = extraer_datos_base(ws)
    assert datos.tiendas[0].tienda == 10081
    assert [(a.codigo, a.fila) for a in datos.articulos] == [(1, 3), (3, 5)]


def test_acepta_eleccion_sin_tilde_y_minusculas(tmp_path):
    ws = hoja(crear_base(tmp_path / "b.xlsx", [(1, 2, 4)], [(10081, {0: (1, 1)})]))
    ws.cell(2, 7).value = " eleccion "
    assert len(extraer_datos_base(ws).tiendas) == 1


def test_ignora_columnas_vacias_con_formato_al_final(tmp_path):
    ws = hoja(crear_base(tmp_path / "b.xlsx", [(1, 2, 4)], [(10081, {0: (1, 1)})]))
    ws.cell(2, 20).number_format = "0.00"
    assert len(extraer_datos_base(ws).tiendas) == 1


# 2.3
def _base_2_tiendas(tmp_path):
    return hoja(crear_base(tmp_path / "b.xlsx", [(1, 2, 4), (2, 3, 5)], [(10081, {0: (1, 1)}), (10082, {1: (2, 1)})]))


def test_error_encabezado_bloque(tmp_path):
    ws = _base_2_tiendas(tmp_path)
    ws["G2"] = "TOTAL"
    with pytest.raises(ErrorValidacion, match=r"Columna G.*ELECCIÓN.*TOTAL"):
        extraer_datos_base(ws)


def test_error_codigo_distinto_en_bloque(tmp_path):
    ws = _base_2_tiendas(tmp_path)
    ws["I1"] = 99999
    with pytest.raises(ErrorValidacion, match=r"Columnas H-J"):
        extraer_datos_base(ws)


def test_error_tienda_repetida(tmp_path):
    ws = _base_2_tiendas(tmp_path)
    for col in ("H1", "I1", "J1"):
        ws[col] = 10081
    with pytest.raises(ErrorValidacion, match=r"tienda 10081 está repetida"):
        extraer_datos_base(ws)


def test_error_columnas_no_multiplo_de_3(tmp_path):
    ws = _base_2_tiendas(tmp_path)
    ws["K1"], ws["K2"] = 10083, "BULTOS"
    with pytest.raises(ErrorValidacion, match=r"múltiplo de 3"):
        extraer_datos_base(ws)


@pytest.mark.parametrize("valor", [-1, "abc"])
@pytest.mark.parametrize("celda,nombre", [("E3", "BULTOS"), ("F3", "LINEAS")])
def test_error_bultos_lineas_no_validos(tmp_path, celda, nombre, valor):
    ws = _base_2_tiendas(tmp_path)
    ws[celda] = valor
    with pytest.raises(ErrorValidacion, match=rf"Fila 3, columna {celda[0]} \(tienda 10081\).*{nombre}"):
        extraer_datos_base(ws)


def test_decimales_y_texto_numerico_admitidos(tmp_path):
    ws = _base_2_tiendas(tmp_path)
    ws["E3"], ws["F3"] = 2.5, "3"
    datos = extraer_datos_base(ws)
    assert datos.tiendas[0].bultos[0] == Fraction(5, 2) and datos.tiendas[0].lineas[0] == 3


# 2.4
@pytest.mark.parametrize("peso,volumen", [(5, 0), (0, 5), (None, 5), (5, "x"), (-1, 5)])
def test_articulo_sin_peso_o_volumen_valido(tmp_path, peso, volumen):
    ws = _base_2_tiendas(tmp_path)
    ws["B3"], ws["C3"] = peso, volumen
    datos = extraer_datos_base(ws)
    assert not datos.articulos[0].valido and datos.articulos[1].valido
    assert len(datos.avisos) == 1 and "Artículo 1 (fila 3)" in datos.avisos[0]


# 2.5
def test_tabla_con_mas_tiendas(tmp_path):
    ruta = crear_tabla(tmp_path / "t.xlsx", [(1, *PALET), (10081, *PALET), (10082, *CARRO), (5, *CARRO)])
    medios = leer_tabla_medios(ruta, [10081, 10082])
    assert set(medios) == {10081, 10082}
    assert medios[10082].tipo == "CARRO" and medios[10082].peso_max == 950


def test_tabla_tiendas_que_faltan(tmp_path):
    ruta = crear_tabla(tmp_path / "t.xlsx", [(10082, *PALET)])
    with pytest.raises(ErrorValidacion, match=r"Faltan 2 tiendas.*10081, 10949"):
        leer_tabla_medios(ruta, [10081, 10082, 10949])


def test_tabla_tipo_desconocido(tmp_path):
    ruta = crear_tabla(tmp_path / "t.xlsx", [(10081, "CAMION", 875, 1300)])
    with pytest.raises(ErrorValidacion, match=r"Tienda 10081.*'CAMION'"):
        leer_tabla_medios(ruta, [10081])


def test_tabla_tipo_sin_distinguir_mayusculas(tmp_path):
    ruta = crear_tabla(tmp_path / "t.xlsx", [("10081", " palet ", 875, 1300)])
    assert leer_tabla_medios(ruta, [10081])[10081].tipo == "PALET"


@pytest.mark.parametrize("col", [3, 4, 5])
def test_tabla_valores_no_validos(tmp_path, col):
    ruta = crear_tabla(tmp_path / "t.xlsx", [(10081, *PALET)])
    wb = openpyxl.load_workbook(ruta)
    wb.active.cell(2, col).value = 0
    with pytest.raises(ErrorValidacion, match=r"Tienda 10081.*mayor que 0"):
        extraer_medios(wb.active, [10081])


def test_tabla_tienda_duplicada_con_datos_distintos(tmp_path):
    ruta = crear_tabla(tmp_path / "t.xlsx", [(10081, *PALET), (10081, *CARRO)])
    with pytest.raises(ErrorValidacion, match=r"10081 aparece con datos distintos"):
        leer_tabla_medios(ruta, [10081])


def test_fichero_inexistente_o_no_excel(tmp_path):
    with pytest.raises(ErrorValidacion, match="No se encuentra"):
        leer_fichero_base(tmp_path / "no_existe.xlsx")
    falso = tmp_path / "falso.xlsx"
    falso.write_text("hola")
    with pytest.raises(ErrorValidacion, match="no es un Excel"):
        leer_tabla_medios(falso, [1])
