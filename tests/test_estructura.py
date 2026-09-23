import openpyxl

import pruebaestandar
from pruebaestandar import escritura, lectura, modelo, proceso, seleccion, verificacion


def test_paquete_importable():
    assert pruebaestandar.__doc__
    for modulo in (escritura, lectura, modelo, proceso, seleccion, verificacion):
        assert modulo.__name__.startswith("pruebaestandar.")


def test_generador_sintetico_3x2(ficheros_pequenos):
    base, tabla = ficheros_pequenos
    ws = openpyxl.load_workbook(base).active
    assert [ws.cell(1, c).value for c in range(5, 11)] == [10081] * 3 + [10082] * 3
    assert [ws.cell(2, c).value for c in range(5, 8)] == ["BULTOS", "LINEAS", "ELECCIÓN"]
    assert [ws.cell(r, 1).value for r in range(3, 6)] == [111, 222, 333]
    assert ws.cell(3, 5).value == 7 and ws.cell(3, 6).value == 3
    wt = openpyxl.load_workbook(tabla).active
    assert wt.max_row == 4 and wt.cell(3, 2).value == "CARRO"
