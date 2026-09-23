import hashlib
import time

import openpyxl
import pytest

from pruebaestandar import escritura, proceso
from pruebaestandar.lectura import leer_fichero_base, leer_tabla_medios
from pruebaestandar.modelo import ErrorEscritura, ErrorValidacion
from pruebaestandar.verificacion import errores_tienda

from .conftest import CARRO, PALET, crear_base, crear_tabla, datos_aleatorios


def huella(ruta):
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def columnas_eleccion(ws):
    return {c for c in range(7, ws.max_column + 1, 3)}


# 4.2
def test_salida_conserva_todo_salvo_eleccion(ficheros_pequenos):
    base, tabla = ficheros_pequenos
    wb = openpyxl.load_workbook(base)
    wb.active["G5"] = 3  # valor previo en ELECCIÓN de un artículo sin salida en la tienda 10081
    wb.active["E30"] = "nota"
    wb.create_sheet("Otra")["A1"] = "se conserva"
    wb.save(base)
    antes = huella(base)

    resumen = proceso.ejecutar(base, tabla, "2")

    assert resumen.ruta_salida == base.parent / "PRUEBAESTANDAR.xlsx"
    assert huella(base) == antes
    ws_in = openpyxl.load_workbook(base).active
    wb_out = openpyxl.load_workbook(resumen.ruta_salida)
    ws_out = wb_out.active
    assert wb_out.sheetnames == ["Hoja1", "Otra"] and wb_out["Otra"]["A1"].value == "se conserva"
    assert ws_out.column_dimensions["A"].width == ws_in.column_dimensions["A"].width == 14
    elec = columnas_eleccion(ws_out)
    for fila in range(1, ws_in.max_row + 1):
        for col in range(1, ws_in.max_column + 1):
            if not (col in elec and 3 <= fila <= 5):
                assert ws_out.cell(fila, col).value == ws_in.cell(fila, col).value, (fila, col)
    # tienda 10081: 111 (7 bultos/2 días -> máx 4), 222 (2/2 -> 1); 333 sin salida
    assert [ws_out.cell(r, 7).value for r in (3, 4, 5)] == [4, 1, None]
    assert [ws_out.cell(r, 10).value for r in (3, 4, 5)] == [2, None, 5]
    assert all(type(ws_out.cell(r, c).value) in (int, type(None)) for r in (3, 4, 5) for c in elec)


def test_fichero_base_llamado_como_la_salida(tmp_path):
    with pytest.raises(ErrorValidacion, match="no puede llamarse PRUEBAESTANDAR.xlsx"):
        escritura.ruta_salida(tmp_path / "pruebaestandar.XLSX")


# 4.3
def test_extremo_a_extremo_20x5(tmp_path):
    articulos, tiendas, tabla = datos_aleatorios(20, 5, semilla=3)
    articulos[4] = (articulos[4][0], 3.0, 0)  # sin volumen -> aviso
    tiendas[2] = (tiendas[2][0], {})  # tienda sin salidas
    base = crear_base(tmp_path / "base.xlsx", articulos, tiendas)
    ruta_tabla = crear_tabla(tmp_path / "tabla.xlsx", tabla)

    resumen = proceso.ejecutar(base, ruta_tabla, "3")

    assert resumen.tiendas_procesadas == 5
    assert resumen.tiendas_sin_seleccion == 1
    assert len(resumen.avisos) == 1 and str(articulos[4][0]) in resumen.avisos[0]
    assert "Tiendas procesadas: 5" in resumen.texto()
    _, datos = leer_fichero_base(resumen.ruta_salida)
    medios = leer_tabla_medios(ruta_tabla, [d.tienda for d in datos.tiendas])
    ws = openpyxl.load_workbook(resumen.ruta_salida).active
    total = 0
    for d in datos.tiendas:
        sel = {i: ws.cell(a.fila, d.col_eleccion).value for i, a in enumerate(datos.articulos)}
        sel = {i: q for i, q in sel.items() if q is not None}
        total += sum(sel.values())
        assert errores_tienda(datos.articulos, d.bultos, medios[d.tienda], 3, sel) == []
        assert 4 not in sel
    assert total == resumen.total_cajas > 0


def test_errores_de_validacion_no_generan_salida(tmp_path):
    base = crear_base(tmp_path / "base.xlsx", [(1, 1, 2)], [(10081, {0: (1, 1)}), (10949, {0: (1, 1)})])
    tabla = crear_tabla(tmp_path / "tabla.xlsx", [(1, *PALET)])
    with pytest.raises(ErrorValidacion, match="10081, 10949"):
        proceso.ejecutar(base, tabla, "5")
    with pytest.raises(ErrorValidacion, match="entero mayor que 0"):
        proceso.ejecutar(base, tabla, "cinco")
    assert not (tmp_path / "PRUEBAESTANDAR.xlsx").exists()


def test_verificacion_fallida_no_guarda(ficheros_pequenos, monkeypatch):
    base, tabla = ficheros_pequenos
    monkeypatch.setattr(proceso, "seleccionar", lambda *a, **k: {0: 999})
    with pytest.raises(ErrorValidacion, match="Tienda 10081"):
        proceso.ejecutar(base, tabla, "1")
    assert not (base.parent / "PRUEBAESTANDAR.xlsx").exists()


# 4.4
def test_reintento_de_guardado_sin_recalcular(ficheros_pequenos, monkeypatch):
    base, tabla = ficheros_pequenos
    calculo = proceso.calcular(base, tabla, "2")
    guardar_original = calculo.libro.save
    intentos = []

    def save_falla_una_vez(destino):
        intentos.append(destino)
        if len(intentos) == 1:
            raise PermissionError("abierto en Excel")
        guardar_original(destino)

    monkeypatch.setattr(calculo.libro, "save", save_falla_una_vez)
    monkeypatch.setattr(proceso, "seleccionar", lambda *a, **k: pytest.fail("no debe recalcular"))
    with pytest.raises(ErrorEscritura, match="ciérralo y pulsa Reintentar"):
        calculo.guardar()
    assert calculo.guardar().ruta_salida.exists()
    assert len(intentos) == 2


# 4.5
@pytest.mark.lento
def test_rendimiento_650x180(tmp_path):
    articulos, tiendas, tabla = datos_aleatorios(650, 180, semilla=11)
    base = crear_base(tmp_path / "base.xlsx", articulos, tiendas)
    ruta_tabla = crear_tabla(tmp_path / "tabla.xlsx", tabla + [(900_000 + i, *CARRO) for i in range(2300)])
    inicio = time.perf_counter()
    resumen = proceso.ejecutar(base, ruta_tabla, "5")
    duracion = time.perf_counter() - inicio
    print(f"650x180: {duracion:.1f} s, {resumen.total_cajas} cajas")
    assert resumen.tiendas_procesadas == 180
    assert duracion < 60
