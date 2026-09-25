import openpyxl
import pytest
import xlrd

from pruebaestandar import proceso, reparto
from pruebaestandar.lectura import leer_fichero_base
from pruebaestandar.modelo import DatosBase, DemandaTienda, ErrorEscritura, ErrorValidacion
from pruebaestandar.verificacion import verificar_reparto

from .conftest import PALET, articulo, crear_base, crear_tabla, datos_aleatorios

BORDE_FINO = (1, 1, 1, 1)


def abrir(ruta):
    libro = xlrd.open_workbook(str(ruta), formatting_info=True)
    return libro, libro.sheet_by_name("Hoja1")


def estilo(libro, hoja, fila, col):
    xf = libro.xf_list[hoja.cell(fila, col).xf_index]
    fuente = libro.font_list[xf.font_index]
    b = xf.border
    return dict(
        formato=libro.format_map[xf.format_key].format_str,
        fuente=fuente.name,
        puntos=fuente.height / 20,
        negrita=bool(fuente.bold),
        color=fuente.colour_index,
        bordes=(b.left_line_style, b.right_line_style, b.top_line_style, b.bottom_line_style),
        relleno=(xf.background.fill_pattern, xf.background.pattern_colour_index),
        alineacion=xf.alignment.hor_align,
    )


def datos_ejemplo():
    arts = [articulo(0, 1, 1, codigo=2101538, unidades_caja=18), articulo(1, 1, 1, codigo=2102526, unidades_caja=6)]
    tiendas = [DemandaTienda(81, 7, [], []), DemandaTienda(82, 10, [], [])]
    return DatosBase(arts, tiendas), {81: {1: 2}, 82: {0: 3}}


# 3.2
def test_unidades_por_tienda_y_celdas_vacias(tmp_path):
    datos, selecciones = datos_ejemplo()
    unidades = reparto.calcular_unidades(datos, selecciones)
    assert unidades == [[None, 54], [12, None]]
    ruta = tmp_path / "REPARTOESTANDAR.xls"
    reparto.construir_reparto(datos, unidades).save(str(ruta))
    _, hoja = abrir(ruta)
    assert hoja.nrows == 3 and hoja.ncols == 3
    assert hoja.row_values(0) == ["Código", 81, 82]
    assert hoja.row_values(1) == [2101538, "", 54]
    assert hoja.row_values(2) == [2102526, 12, ""]
    assert [hoja.cell(0, c).ctype for c in range(3)] == [xlrd.XL_CELL_TEXT, xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_NUMBER]
    assert hoja.cell(1, 1).ctype == xlrd.XL_CELL_BLANK and hoja.cell(1, 0).ctype == xlrd.XL_CELL_NUMBER


# 3.3
def test_formato_como_la_plantilla(tmp_path):
    datos, selecciones = datos_ejemplo()
    ruta = tmp_path / "REPARTOESTANDAR.xls"
    reparto.construir_reparto(datos, reparto.calcular_unidades(datos, selecciones)).save(str(ruta))
    libro, hoja = abrir(ruta)
    assert ruta.read_bytes()[:8] == bytes.fromhex("d0cf11e0a1b11ae1")  # contenedor OLE2 de Excel 97-2003
    assert libro.biff_version == 80 and libro.sheet_names() == ["Hoja1", "Hoja2", "Hoja3"]
    assert libro.sheet_by_name("Hoja2").nrows == libro.sheet_by_name("Hoja3").nrows == 0
    assert estilo(libro, hoja, 0, 0) == dict(
        formato="General", fuente="Arial", puntos=12, negrita=True, color=8, bordes=BORDE_FINO, relleno=(1, 44), alineacion=2
    )
    assert estilo(libro, hoja, 0, 1) == dict(
        formato="General", fuente="Arial", puntos=10, negrita=True, color=12, bordes=BORDE_FINO, relleno=(0, 64), alineacion=0
    )
    assert estilo(libro, hoja, 1, 0) == dict(
        formato="0", fuente="Arial", puntos=10, negrita=False, color=32767, bordes=BORDE_FINO, relleno=(0, 64), alineacion=0
    )
    for celda in ((1, 1), (1, 2)):  # vacía y con valor usan el mismo estilo
        assert estilo(libro, hoja, *celda) == dict(
            formato="General", fuente="Arial", puntos=12, negrita=True, color=32767, bordes=BORDE_FINO, relleno=(0, 64), alineacion=0
        )
    assert (hoja.colinfo_map[0].width, hoja.colinfo_map[1].width, hoja.colinfo_map[2].width) == (2304, 1024, 1024)
    assert {hoja.rowinfo_map[r].height for r in range(3)} == {315}


# 3.4
def test_mas_de_255_tiendas_no_genera_nada(tmp_path):
    tiendas = [(t, {0: (1, 1)}) for t in range(1, 301)]
    base = crear_base(tmp_path / "base.xlsx", [(1, 1, 2)], tiendas)
    tabla = crear_tabla(tmp_path / "tabla.xlsx", [(t, *PALET) for t in range(1, 301)])
    with pytest.raises(ErrorValidacion, match="300 tiendas.*como máximo 255"):
        proceso.ejecutar(base, tabla, "5")
    assert not list(tmp_path.glob("PRUEBAESTANDAR*")) and not list(tmp_path.glob("REPARTOESTANDAR*"))


def test_255_tiendas_si_caben(tmp_path):
    tiendas = [(t, {0: (1, 1)}) for t in range(1, 256)]
    base = crear_base(tmp_path / "base.xlsx", [(7, 1, 2, 4)], tiendas)
    tabla = crear_tabla(tmp_path / "tabla.xlsx", [(t, *PALET) for t in range(1, 256)])
    resumen = proceso.ejecutar(base, tabla, "5")
    _, hoja = abrir(resumen.ruta_reparto)
    assert hoja.ncols == 256 and hoja.cell(0, 255).value == 255 and hoja.cell(1, 255).value == 4


# 4.1
def test_verificacion_detecta_reparto_manipulado():
    datos, selecciones = datos_ejemplo()
    unidades = reparto.calcular_unidades(datos, selecciones)
    verificar_reparto(datos, selecciones, unidades)
    unidades[0][1] = 53
    with pytest.raises(ErrorValidacion, match=r"Tienda 82, artículo 2101538: el reparto tiene 53 unidades y deberían ser 54"):
        verificar_reparto(datos, selecciones, unidades)


def test_verificacion_detecta_forma_incorrecta():
    datos, selecciones = datos_ejemplo()
    with pytest.raises(ErrorValidacion, match="una fila por artículo"):
        verificar_reparto(datos, selecciones, [[None, 54]])


def test_reparto_manipulado_no_guarda_ningun_fichero(ficheros_pequenos, monkeypatch):
    base, tabla = ficheros_pequenos
    original = reparto.calcular_unidades
    monkeypatch.setattr(reparto, "calcular_unidades", lambda d, s: [[(v or 0) + 1 for v in f] for f in original(d, s)])
    with pytest.raises(ErrorValidacion, match="reparto en unidades"):
        proceso.ejecutar(base, tabla, "2")
    assert not (base.parent / "PRUEBAESTANDAR.xlsx").exists() and not (base.parent / "REPARTOESTANDAR.xls").exists()


# 4.2
def test_extremo_a_extremo_reparto_igual_a_eleccion_por_unidades(tmp_path):
    articulos, tiendas, tabla = datos_aleatorios(40, 12, semilla=5)
    base = crear_base(tmp_path / "base.xlsx", articulos, tiendas)
    resumen = proceso.ejecutar(base, crear_tabla(tmp_path / "tabla.xlsx", tabla), "4")
    assert resumen.ruta_reparto == tmp_path / "REPARTOESTANDAR.xls"
    assert "REPARTOESTANDAR.xls" in resumen.texto() and f"Total de unidades: {resumen.total_unidades}" in resumen.texto()
    _, datos = leer_fichero_base(resumen.ruta_salida)
    ws = openpyxl.load_workbook(resumen.ruta_salida).active
    _, hoja = abrir(resumen.ruta_reparto)
    assert hoja.nrows == 41 and hoja.ncols == 13
    total = 0
    for c, d in enumerate(datos.tiendas, start=1):
        assert hoja.cell(0, c).value == d.tienda
        for r, art in enumerate(datos.articulos, start=1):
            assert hoja.cell(r, 0).value == art.codigo
            cajas = ws.cell(art.fila, d.col_eleccion).value
            esperado = "" if cajas is None else cajas * art.unidades_caja
            assert hoja.cell(r, c).value == esperado, (d.tienda, art.codigo)
            total += 0 if cajas is None else cajas * art.unidades_caja
    assert total == resumen.total_unidades > 0


def test_reintento_con_reparto_bloqueado(ficheros_pequenos, monkeypatch):
    base, tabla = ficheros_pequenos
    calculo = proceso.calcular(base, tabla, "2")
    guardar_original = calculo.libro_reparto.save
    intentos = []

    def save_falla_una_vez(destino):
        intentos.append(destino)
        if len(intentos) == 1:
            raise PermissionError("abierto en Excel")
        guardar_original(destino)

    monkeypatch.setattr(calculo.libro_reparto, "save", save_falla_una_vez)
    monkeypatch.setattr(proceso, "seleccionar", lambda *a, **k: pytest.fail("no debe recalcular"))
    with pytest.raises(ErrorEscritura, match="REPARTOESTANDAR.xls.*ciérralo y pulsa Reintentar"):
        calculo.guardar()
    resumen = calculo.guardar()
    assert resumen.ruta_salida.exists() and resumen.ruta_reparto.exists() and len(intentos) == 2
